import os
import re
import io
import json
from flask import Flask, request, send_file, render_template, jsonify, redirect, url_for
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from google import genai
from google.genai import types
from dotenv import load_dotenv
 


app = Flask(__name__, static_folder='public/static', static_url_path='/static')


# Load environment variables
load_dotenv()

# Configure APIs safely
GEMINI_KEY = os.getenv("GEMINI_API_KEY")

if GEMINI_KEY:
    gemini_client = genai.Client(
        api_key=GEMINI_KEY,
        http_options=types.HttpOptions(
            timeout=18_000,
            retry_options=types.HttpRetryOptions(attempts=1),
        ),
    )
else:
    gemini_client = None
    print("Warning: GEMINI_API_KEY not found.")

GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite")
MAX_SLIDES = 12
ALLOWED_TEMPLATES = {"modern", "minimal", "creative", "dark"}


class AIConfigurationError(RuntimeError):
    """Raised when required Gemini configuration is missing."""


def ask_ai(prompt, json_output=False):
    """Generate text with the configured Gemini model."""
    if not gemini_client:
        raise AIConfigurationError("Gemini API key is missing. Add GEMINI_API_KEY to your .env file and restart the app.")
    try:
        response_config = {"response_mime_type": "application/json"} if json_output else None
        if json_output:
            response_config["max_output_tokens"] = 1800
            response_config["thinking_config"] = {"thinking_level": "minimal"}
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config=response_config,
        )
        text = getattr(response, "text", "")
        if text and text.strip():
            return text.strip()
        raise RuntimeError("Gemini returned an empty response.")
    except RuntimeError:
        raise
    except Exception as exc:
        status = getattr(exc, "code", None)
        app.logger.warning(
            "Gemini request failed: exception=%s provider_status=%s",
            type(exc).__name__,
            status if isinstance(status, int) else "unknown",
        )
        if isinstance(exc, TimeoutError) or "timeout" in type(exc).__name__.lower():
            message = "Gemini took longer than 18 seconds to respond. Retry with fewer slides or try again shortly."
        elif status in (400, 404):
            message = "Gemini rejected the request or model. Check GEMINI_MODEL and try again."
        elif status in (401, 403):
            message = "Gemini rejected the API key or this key lacks API access. Check GEMINI_API_KEY in .env."
        elif status == 429:
            message = "Gemini API quota or rate limit reached. Check your Google AI Studio quota and retry later."
        else:
            message = "Gemini request failed. Check the server log for the provider status code."
        raise RuntimeError(message) from None


# Slide creation functions (unchanged from your code)
def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return RGBColor(*tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4)))

def get_template_styles(template_name, theme_colors):
    """Define styles for different templates."""
    primary = theme_colors[0]
    secondary = theme_colors[1]
    white = RGBColor(255, 255, 255)
    dark = RGBColor(20, 20, 25)

    styles = {
        'modern': {
            'bg': white,
            'title_color': hex_to_rgb(primary),
            'text_color': RGBColor(44, 62, 80),
            'accent': hex_to_rgb(secondary),
            'shape': MSO_SHAPE.RECTANGLE,
            'cover_bg': hex_to_rgb(primary),
            'cover_title_color': white,
            'outline_bg': RGBColor(240, 244, 250),
            'outline_card_bg': white,
            'outline_card_text': RGBColor(30, 41, 59),
            'outline_shape': MSO_SHAPE.ROUNDED_RECTANGLE,
            'watermark': RGBColor(148, 163, 184),
        },
        'minimal': {
            'bg': white,
            'title_color': RGBColor(31, 41, 55),
            'text_color': RGBColor(50, 50, 50),
            'accent': RGBColor(203, 213, 225),
            'shape': None,
            'cover_bg': RGBColor(248, 250, 252),
            'cover_title_color': RGBColor(15, 23, 42),
            'outline_bg': white,
            'outline_card_bg': white,
            'outline_card_text': RGBColor(31, 41, 55),
            'outline_shape': MSO_SHAPE.ROUNDED_RECTANGLE,
            'watermark': RGBColor(148, 163, 184),
        },
        'creative': {
            'bg': RGBColor(245, 245, 250),
            'title_color': hex_to_rgb(primary),
            'text_color': RGBColor(33, 33, 33),
            'accent': hex_to_rgb(secondary),
            'shape': MSO_SHAPE.ROUNDED_RECTANGLE,
            'cover_bg': hex_to_rgb(primary),
            'cover_title_color': white,
            'outline_bg': RGBColor(245, 245, 250),
            'outline_card_bg': white,
            'outline_card_text': RGBColor(33, 33, 33),
            'outline_shape': MSO_SHAPE.ROUNDED_RECTANGLE,
            'watermark': RGBColor(120, 120, 135),
        },
        'dark': {
            'bg': dark,
            'title_color': white,
            'text_color': RGBColor(200, 200, 200),
            'accent': hex_to_rgb(secondary),
            'shape': MSO_SHAPE.RECTANGLE,
            'cover_bg': dark,
            'cover_title_color': white,
            'outline_bg': dark,
            'outline_card_bg': RGBColor(40, 40, 50),
            'outline_card_text': RGBColor(235, 235, 240),
            'outline_shape': MSO_SHAPE.ROUNDED_RECTANGLE,
            'watermark': RGBColor(148, 148, 160),
        }
    }
    return styles.get(template_name, styles['modern'])

def generate_deck_content(title, topic, language, num_slides):
    """Generate the complete deck plan in one model request."""
    prompt = f"""
Create content for an editable business presentation.
Title: {title}
Topic: {topic}
Language: {language}
Return exactly {num_slides} content slides. Keep slide titles distinct and concise.
Each slide must have 4 or 5 concise, factual bullets, each 8 words maximum.
Keep every point directly relevant to the topic. Do not add a key insight, summary sentence, conclusion, or any text outside the bullets.
Return JSON with this exact shape and no extra keys:
{{
  "primary_color": "#1E3A8A",
  "secondary_color": "#F59E0B",
  "slides": [
    {{"title": "Slide title", "bullets": ["Point one", "Point two", "Point three", "Point four"]}}
  ]
}}
Use valid six-digit hex colors. Do not include markdown fences.
"""
    raw = ask_ai(prompt, json_output=True)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Gemini returned malformed presentation data. Please retry.") from exc

    slides = data.get("slides") if isinstance(data, dict) else None
    if not isinstance(slides, list) or len(slides) != num_slides:
        raise RuntimeError("Gemini returned an incomplete slide plan. Please retry.")

    clean_slides = []
    for slide in slides:
        if not isinstance(slide, dict):
            raise RuntimeError("Gemini returned an invalid slide. Please retry.")
        slide_title = slide.get("title")
        bullets = slide.get("bullets")
        if (not isinstance(slide_title, str) or not slide_title.strip()
                or not isinstance(bullets, list) or not 4 <= len(bullets) <= 5
                or not all(isinstance(point, str) and point.strip() for point in bullets)):
            raise RuntimeError("Gemini returned an invalid slide. Please retry.")
        body = "\n".join(f"• {point.strip()}" for point in bullets)
        clean_slides.append((slide_title.strip(), body))

    colors = [data.get("primary_color", ""), data.get("secondary_color", "")]
    if not all(isinstance(color, str) and re.fullmatch(r"#[0-9a-fA-F]{6}", color) for color in colors):
        colors = ["#0D47A1", "#FF9800"]
    return colors, clean_slides

def add_welcome_slide(prs, title_text, theme, template_style):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    set_slide_background(slide, template_style['cover_bg'])
    for shape in slide.placeholders:
        sp = shape.element
        sp.getparent().remove(sp)
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(2)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(48)
    p.font.color.rgb = template_style['cover_title_color']
    p.alignment = PP_ALIGN.CENTER
    add_decorative_elements(slide, theme, template_style)

def add_content_page(prs, sections, topic, theme, template_style):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for shape in slide.placeholders:
        sp = shape.element
        sp.getparent().remove(sp)
    set_slide_background(slide, template_style['outline_bg'])
    left = Inches(0.5)
    top = Inches(0.5)
    width = Inches(9)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = f"Presentation Outline: {topic}"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = template_style['title_color']
    p.alignment = PP_ALIGN.LEFT
    columns = 3
    rows = (len(sections) + columns - 1) // columns
    card_height = min(1.25, 4.55 / max(rows, 1))
    row_gap = (4.55 - rows * card_height) / max(rows - 1, 1)
    for i, section in enumerate(sections):
        row = i // columns
        col = i % columns
        left_pos = Inches(0.5 + col * 3)
        top_pos = Inches(1.65 + row * (card_height + row_gap))
        shape = slide.shapes.add_shape(
            template_style['outline_shape'],
            left_pos, top_pos, Inches(2.8), Inches(card_height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = template_style['outline_card_bg']
        shape.line.color.rgb = template_style['accent']
        shape.line.width = Pt(1.5)
        tf = shape.text_frame
        p = tf.add_paragraph()
        p.text = f"{i+1}. {section}"
        p.font.size = Pt(14)
        p.font.color.rgb = template_style['outline_card_text']
        p.alignment = PP_ALIGN.CENTER
    add_watermark(slide, template_style)

def add_content_slide(prs, title_text, content_text, topic, template_style, section):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for shape in slide.placeholders:
        sp = shape.element
        sp.getparent().remove(sp)
        
    set_slide_background(slide, template_style['bg'])
    
    # Title
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(32)
    p.font.color.rgb = template_style['title_color']
    p.alignment = PP_ALIGN.LEFT

    # Accent Line
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left, Inches(1.1), width, Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = template_style['accent']
    line.line.visible = False

    # Content Box (Full Width Minimalist)
    content_width = Inches(8.8)
    content_left = Inches(0.6)
    
    if template_style['shape']:
        content_box = slide.shapes.add_shape(
            template_style['shape'],
            content_left, Inches(1.65), content_width, Inches(4.7))
        content_box.fill.solid()
        if template_style['bg'] == RGBColor(20, 20, 25): # Dark mode
            content_box.fill.fore_color.rgb = RGBColor(40, 40, 50)
        else:
            content_box.fill.fore_color.rgb = RGBColor(248, 249, 250)
        content_box.line.color.rgb = template_style['accent']
        content_box.line.width = Pt(1)
        tf = content_box.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE
    else:
        txBox = slide.shapes.add_textbox(content_left, Inches(1.65), content_width, Inches(4.7))
        tf = txBox.text_frame
        tf.auto_size = MSO_AUTO_SIZE.NONE

    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.2)
    tf.clear()
    for index, point in enumerate(line.strip() for line in content_text.splitlines() if line.strip()):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = point
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        p.font.size = Pt(24)
        p.font.color.rgb = template_style['text_color']
    
    add_watermark(slide, template_style)

def add_thank_you_slide(prs, theme, template_style):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for shape in slide.placeholders:
        sp = shape.element
        sp.getparent().remove(sp)
    set_slide_background(slide, template_style['cover_bg'])
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(2)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "Thank You"
    p.font.bold = True
    p.font.size = Pt(60)
    p.font.color.rgb = template_style['cover_title_color']
    p.alignment = PP_ALIGN.CENTER
    add_decorative_elements(slide, theme, template_style)

def add_decorative_elements(slide, theme, template_style):
    for i in range(3):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(0.5 + i*3), Inches(6.5), Inches(0.5), Inches(0.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = template_style['accent']
        shape.line.color.rgb = template_style['accent']

def add_watermark(slide, template_style):
    left = Inches(7.5)
    top = Inches(6.8)
    width = Inches(2.5)
    height = Inches(0.5)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "Generated by Subham Khandual | Sayraa AI PPT Generator"
    p.font.size = Pt(9)
    p.font.color.rgb = template_style['watermark']
    p.alignment = PP_ALIGN.RIGHT

def create_presentation(title, topic, language, num_slides, template_name):
    prs = Presentation()
    # Set document properties
    core_props = prs.core_properties
    core_props.author = "Subham Khandual"
    core_props.title = title
    
    theme, slide_content = generate_deck_content(title, topic, language, num_slides)
    style = get_template_styles(template_name, theme)
    sections = [slide_title for slide_title, _ in slide_content]
    
    add_welcome_slide(prs, title, theme, style)
    add_content_page(prs, sections, topic, theme, style)
    
    for section, (_, content) in zip(sections, slide_content):
        add_content_slide(prs, section, content, topic, style, section)
        
    add_thank_you_slide(prs, theme, style)
    
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    
    return buffer

# Flask routes
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/favicon.ico')
def favicon():
    return redirect(url_for('static', filename='logo-192.png'))

@app.route('/manifest.json')
def manifest():
    return redirect(url_for('static', filename='manifest.json'))

@app.route('/sitemap.xml')
def sitemap():
    import datetime
    today = datetime.date.today().isoformat()
    sitemap_xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
  <url>
    <loc>{request.host_url}</loc>
    <lastmod>{today}</lastmod>
    <changefreq>weekly</changefreq>
    <priority>1.0</priority>
  </url>
</urlset>"""
    return app.response_class(sitemap_xml, mimetype='application/xml')

@app.route('/robots.txt')
def robots():
    robots_txt = f"User-agent: *\nAllow: /\nSitemap: {request.host_url}sitemap.xml\n"
    return app.response_class(robots_txt, mimetype='text/plain')

def cleanup_old_files():
    try:
        import tempfile
        import time
        temp_dir = os.path.join(tempfile.gettempdir(), 'sayraa_ppt')
        if os.path.exists(temp_dir):
            now = time.time()
            for f in os.listdir(temp_dir):
                fp = os.path.join(temp_dir, f)
                if os.path.isfile(fp):
                    # Delete files older than 10 minutes (600 seconds)
                    if now - os.path.getmtime(fp) > 600:
                        try:
                            os.remove(fp)
                        except Exception:
                            pass
    except Exception as e:
        print(f"Error cleaning up old files: {e}")

@app.route('/generate', methods=['POST'])
def generate_ppt():
    cleanup_old_files()
    if request.is_json:
        data = request.get_json()
    else:
        data = request.form
        
    title = data.get('title')
    topic = data.get('topic')
    language = data.get('language', 'English')
    template = data.get('template', 'modern')
    try:
        num_slides = int(data.get('num_slides', 8))
    except (TypeError, ValueError):
        num_slides = 0

    # Helper function to create error response with cookie
    def make_error_response(err_msg, code=500):
        import urllib.parse
        safe_msg = err_msg if code != 500 else 'Presentation generation failed. Check the server log and try again.'
        encoded_msg = urllib.parse.quote(safe_msg)
        response = jsonify({'error': safe_msg})
        response.status_code = code
        response.set_cookie('ppt_error', encoded_msg, path='/')
        return response

    if not isinstance(title, str) or not isinstance(topic, str) or not title.strip() or not topic.strip():
        return make_error_response('Title and topic are required', 400)
    title, topic = title.strip(), topic.strip()
    if len(title) > 120 or len(topic) > 1200:
        return make_error_response('Title must be at most 120 characters and topic at most 1,200 characters.', 400)
    if not isinstance(language, str) or len(language) > 40:
        return make_error_response('Choose a supported presentation language.', 400)
    if template not in ALLOWED_TEMPLATES:
        return make_error_response('Choose a supported presentation template.', 400)
    if not 3 <= num_slides <= MAX_SLIDES:
        return make_error_response(f'Choose between 3 and {MAX_SLIDES} content slides.', 400)

    try:
        buffer = create_presentation(title, topic, language, num_slides, template)
        filename = re.sub(r'[^\w.-]+', '_', title, flags=re.UNICODE).strip('._') or 'presentation'
        response = send_file(
            buffer,
            as_attachment=True,
            download_name=f"{filename}_Presentation.pptx",
            mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )
        response.set_cookie('ppt_downloaded', 'true', path='/')
        return response
    except AIConfigurationError as e:
        app.logger.error("Gemini configuration missing")
        return make_error_response(str(e), 503)
    except RuntimeError as e:
        app.logger.error("Presentation generation could not complete: %s", str(e))
        return make_error_response(str(e), 502)
    except Exception:
        app.logger.exception("Presentation generation failed")
        return make_error_response('Presentation generation failed. Check your AI provider configuration and try again.', 500)

if __name__ == '__main__':
    app.run(debug=os.getenv('FLASK_DEBUG', '0') == '1', port=int(os.getenv('PORT', '5000')))
